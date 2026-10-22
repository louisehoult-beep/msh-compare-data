#!/usr/bin/env python3
"""ICB pack data builder (dermatology, emollients). Builds data/icb-pack/ for the Hub tool
app/icb-pack.js: index.json (ranking inputs for every ICB and brand, national series, RTT,
IMD, licences) plus one icb-<CODE>.json per ICB (monthly brand and market items, PCN rows).

Spec: 02-Elevate-and-Thrive/Process flows for all brands/icb-pack-dermatology-spec-2026-10-22.md
Check: python3 test_icb_pack.py

Stages
  fetch   download the open sources into a cache folder (default tmp/icb-pack-cache,
          git-ignored). Needs the network. Skipped when the files are already there
          unless --refresh is given.
  build   compute every figure and write data/icb-pack/.

Nothing is invented. Where a figure cannot be had from a named source it is null and the
tool prints "Not verified" (HUB-VERIFICATION-STANDARD).

Brand = first word of the NHSBSA presentation name, for first words with at least
BRAND_FLOOR national items in the 12 months that are not generic words. Market =
emollient preparations: BNF 130201 plus 2122 (appliance-list emollients). The
denominator is OUR definition, shown in the tool: Aneeka's own is not known.
"""
import argparse
import collections
import csv
import datetime
import glob
import html
import io
import json
import math
import os
import re
import statistics
import sys
import urllib.request
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

GP_DIR = os.path.join(ROOT, "data", "gp-prescribing")
UA = {"User-Agent": "msh-compare-data/icb-pack (+elevateandthrive.uk)"}

# Source landing pages. File URLs carry content hashes that change every release, so
# they are resolved from these pages at fetch time, never stored.
RTT_PAGE = "https://www.england.nhs.uk/statistics/statistical-work-areas/rtt-waiting-times/rtt-data-2026-27/"
IMD_API = "https://www.gov.uk/api/content/government/statistics/english-indices-of-deprivation-2025"
# ONS LSOA (2021) to ICB (April 2026) lookup; item id from data.gov.uk dataset 1d16e489-1dc8-443e-9e1c-15de4f961c6d
LSOA_ICB_LOOKUP = "https://open-geography-portalx-ons.hub.arcgis.com/api/download/v1/items/bebf631bb1ac40cf9ec7dc060897ca99/csv?layers=0"
GPREG_PARENT = "https://digital.nhs.uk/data-and-information/publications/statistical/patients-registered-at-a-gp-practice"

LICENCES = [
    ("NHSBSA English Prescribing Dataset (SNOMED)", "Open Government Licence 3.0; reuse statement: NHSBSA Copyright. "
     "https://opendata.nhsbsa.net/dataset/english-prescribing-dataset-epd-with-snomed-code"),
    ("NHS England RTT waiting times (Incomplete Commissioner)", "Open Government Licence v3.0. " + RTT_PAGE),
    ("English Indices of Deprivation 2025, File 13 ICB summaries (MHCLG)", "Open Government Licence v3.0. "
     "https://www.gov.uk/government/statistics/english-indices-of-deprivation-2025"),
    ("IMD 2025 File 7 (LSOA scores and population) with ONS LSOA (2021) to ICB (April 2026) lookup",
     "Open Government Licence v3.0 (IMD); ONS lookup: data.gov.uk dataset licence field unset, site footer says "
     "Open Government Licence v3.0 except where stated. https://www.data.gov.uk/dataset/1d16e489-1dc8-443e-9e1c-15de4f961c6d"),
    ("NHS Digital Patients Registered at a GP Practice, mapping and practice totals", "NHS England terms: "
     "released under the Open Government Licence. " + GPREG_PARENT),
    ("NHS England, items which should not routinely be prescribed in primary care", "Open Government Licence v3.0. "
     "https://www.england.nhs.uk/long-read/items-which-should-not-routinely-be-prescribed-in-primary-care-policy-guidance"),
]

NOT_VERIFIED = "Not verified"
MARKET_PREFIXES = ("130201", "2122")      # BNF presentation-code prefixes


# --------------------------------------------------------------------------
# fetching
# --------------------------------------------------------------------------
def http(url, timeout=120, retries=2):
    last = None
    for _ in range(retries + 1):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r:
                return r.read()
        except Exception as exc:  # noqa: BLE001 - retried then reported
            last = exc
    raise RuntimeError("could not fetch %s: %s" % (url.split("?")[0], last))


def resolve_urls():
    """Current file URLs from the publishers' own pages."""
    urls = {}
    page = http(RTT_PAGE).decode("utf-8", "replace")
    m = re.search(r'href="([^"]*Incomplete-Commissioner-[A-Za-z]{3}\d\d[^"]*\.xlsx)"', page)
    if not m:
        raise RuntimeError("RTT Incomplete Commissioner link not found on %s" % RTT_PAGE)
    urls["rtt"] = html.unescape(m.group(1))
    doc = json.loads(http(IMD_API))
    atts = (doc.get("details") or {}).get("attachments", [])
    m = next((a for a in atts if "Integrated Care Board" in (a.get("title") or "") and a.get("url", "").endswith(".xlsx")), None)
    if not m:
        raise RuntimeError("IMD 2025 File 13 link not found")
    urls["imd"] = m["url"]
    f7 = next((a for a in atts if (a.get("title") or "").startswith("File 7")), None)
    if not f7:
        raise RuntimeError("IMD 2025 File 7 link not found")
    urls["imd7"] = f7["url"]
    parent = http(GPREG_PARENT).decode("utf-8", "replace")
    pubs = sorted(set(re.findall(r'href="([^"]*patients-registered-at-a-gp-practice/[a-z]+-\d{4})"', parent)))
    if not pubs:
        raise RuntimeError("no Patients Registered publication links found")
    months = ["january", "february", "march", "april", "may", "june", "july", "august",
              "september", "october", "november", "december"]

    def key(u):
        mo, yr = u.rsplit("/", 1)[1].split("-")
        return (int(yr), months.index(mo))
    # Newest first; a forthcoming month is listed before it has files, so take the first
    # publication that actually carries the mapping zip.
    for latest in sorted(pubs, key=key, reverse=True):
        if latest.startswith("/"):
            latest = "https://digital.nhs.uk" + latest
        try:
            pub = http(latest).decode("utf-8", "replace")
        except RuntimeError:
            continue
        found = {}
        for tag, pat in (("gpmap", r"gp-reg-pat-prac-map\.zip"), ("gpall", r"gp-reg-pat-prac-all\.zip")):
            m = re.search(r'href="([^"]*%s)"' % pat, pub)
            if m:
                found[tag] = m.group(1)
        if len(found) == 2:
            urls.update(found)
            break
    else:
        raise RuntimeError("no Patients Registered publication with the mapping and practice files")
    urls["gpreg_page"] = latest
    return urls


def fetch_all(cache, refresh=False):
    os.makedirs(cache, exist_ok=True)
    meta_path = os.path.join(cache, "sources.json")
    files = ("rtt-incomplete-commissioner.xlsx", "imd2025-file13-icb.xlsx", "imd2025-file7-lsoa.csv",
             "lsoa21-icb26-lookup.csv", "gp-reg-pat-prac-map.zip", "gp-reg-pat-prac-all.zip")
    need = refresh or not os.path.exists(meta_path) or any(
        not os.path.exists(os.path.join(cache, f)) for f in files)
    if need:
        urls = resolve_urls()
        urls["lsoa"] = LSOA_ICB_LOOKUP
        for tag, fn in (("rtt", "rtt-incomplete-commissioner.xlsx"), ("imd", "imd2025-file13-icb.xlsx"), ("imd7", "imd2025-file7-lsoa.csv"),
                        ("lsoa", "lsoa21-icb26-lookup.csv"),
                        ("gpmap", "gp-reg-pat-prac-map.zip"), ("gpall", "gp-reg-pat-prac-all.zip")):
            with open(os.path.join(cache, fn), "wb") as f:
                f.write(http(urls[tag], timeout=240))
        urls["fetchedOn"] = datetime.date.today().isoformat()
        json.dump(urls, open(meta_path, "w"), indent=1)
    return json.load(open(meta_path))


def fetch_rtt_history(cache, refresh=False, months=13):
    """The last `months` Incomplete Commissioner files, from the two NHS England RTT year pages."""
    os.makedirs(os.path.join(cache, "rtt"), exist_ok=True)
    found = {}
    mon = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
    for page in (RTT_PAGE, RTT_PAGE.replace("2026-27", "2025-26")):
        text = http(page).decode("utf-8", "replace")
        for u in re.findall(r'href="([^"]*Incomplete-Commissioner-[A-Za-z]{3}\d\d[^"]*\.xlsx)"', text):
            m = re.search(r"Commissioner-([A-Za-z]{3})(\d\d)", u)
            found["20%s%02d" % (m.group(2), mon[m.group(1)])] = html.unescape(u)
    wanted = sorted(found)[-months:]
    for p in wanted:
        path = os.path.join(cache, "rtt", "%s.xlsx" % p)
        if refresh or not os.path.exists(path):
            with open(path, "wb") as f:
                f.write(http(found[p], timeout=240))
    return wanted


def load_rtt_history(cache):
    """{period: {icb_code: row}} for dermatology from every cached month."""
    out = {}
    for path in sorted(glob.glob(os.path.join(cache, "rtt", "*.xlsx"))):
        out[os.path.basename(path)[:6]] = load_rtt(cache, path)[0]
    return out



# --------------------------------------------------------------------------
# practice-level pull: every emollient brand, not just one
# --------------------------------------------------------------------------
# First words of presentation names that are generics or formulations, not brands.
GENERIC_WORDS = {"aqueous", "emulsifying", "liquid", "white", "isopropyl", "urea", "yellow", "paraffin",
                 "soft", "hard", "wool", "glycerin", "glycerol", "sodium", "coal", "arachis", "zinc"}
BRAND_FLOOR = 5000        # national 12-month items a first word needs to be offered as a brand
OTHER = "other"


def fetch_practice_level(cache, periods, refresh=False):
    """Practice x brand-word items and cost per month, England, through the NHSBSA datastore
    SQL API (server-side aggregation; ~133,000 rows a month, ~8 seconds)."""
    import refresh_gp_prescribing as g
    path = os.path.join(cache, "epd-practice-brand-months.json")
    have = json.load(open(path)) if os.path.exists(path) and not refresh else {}
    res = dict(g.resources())
    for p in periods:
        if p in have:
            continue
        t = res[p]
        q = ("SELECT PRACTICE_CODE AS pr, SPLIT(BNF_PRESENTATION_NAME,' ')[OFFSET(0)] AS b, "
             "SUM(ITEMS) AS i, SUM(ACTUAL_COST) AS c FROM `{t}` "
             "WHERE SUBSTR(BNF_PRESENTATION_CODE,1,6)='130201' OR SUBSTR(BNF_PRESENTATION_CODE,1,4)='2122' "
             "GROUP BY PRACTICE_CODE, b".format(t=t))
        month = collections.defaultdict(dict)
        for r in g.sql(t, q):
            month[r["pr"]][r["b"]] = [float(r["i"] or 0), float(r["c"] or 0)]
        have[p] = month
        json.dump(have, open(path, "w"))
    return have


def brand_key(word):
    return re.sub(r"[^a-z0-9]", "", (word or "").lower())

def zip_csv(path):
    with zipfile.ZipFile(path) as z:
        name = [n for n in z.namelist() if n.lower().endswith(".csv")][0]
        return list(csv.DictReader(io.StringIO(z.read(name).decode("utf-8-sig"))))


def load_gpreg(cache):
    mp = {}
    for r in zip_csv(os.path.join(cache, "gp-reg-pat-prac-map.zip")):
        mp[r["PRACTICE_CODE"]] = {"pcn": r["PCN_CODE"], "pcn_name": r["PCN_NAME"], "icb": r["ICB_CODE"],
                                  "icb_name": r["ICB_NAME"], "ons_icb": r["ONS_ICB_CODE"], "name": r["PRACTICE_NAME"]}
    size = {}
    for r in zip_csv(os.path.join(cache, "gp-reg-pat-prac-all.zip")):
        if r["TYPE"] == "GP" and r["SEX"] == "ALL" and r["AGE"] == "ALL":
            size[r["CODE"]] = int(r["NUMBER_OF_PATIENTS"])
    return mp, size


def load_rtt(cache, path=None):
    """Dermatology Service (C_330) rows: per ICB from the ICB sheet, England from the
    National sheet. Returns ({icb_code: row, "ENGLAND": row}, period)."""
    import openpyxl
    wb = openpyxl.load_workbook(path or os.path.join(cache, "rtt-incomplete-commissioner.xlsx"),
                                read_only=True, data_only=True)
    need = ["Total number of incomplete pathways", "Total within 18 weeks",
            "Average (median) waiting time (in weeks)", "Total 52 plus weeks", "Total 65 plus weeks",
            "Total 78 plus weeks"]
    out, period = {}, None

    def rows_of(sheet):
        rows = list(wb[sheet].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r[1] in ("ICB Code", "Treatment Function Code"))
        h = {str(c).strip(): i for i, c in enumerate(rows[hi]) if c is not None}
        return rows, hi, h

    def pick(r, h, name):
        return {"name": name, "total": r[h[need[0]]], "w18": r[h[need[1]]], "median": r[h[need[2]]],
                "w52": r[h[need[3]]], "w65": r[h[need[4]]], "w78": r[h[need[5]]]}
    rows, hi, h = rows_of("ICB")
    for r in rows[:12]:
        if r[1] == "Period:":
            period = r[2]
    for r in rows[hi + 1:]:
        if r[3] == "C_330" and r[1] != "-":
            out[r[1]] = pick(r, h, r[2])
    rows, hi, h = rows_of("National")
    for r in rows[hi + 1:]:
        if r[1] == "C_330":
            out["ENGLAND"] = pick(r, h, "England")
    return out, period


def load_imd(cache):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(cache, "imd2025-file13-icb.xlsx"), read_only=True, data_only=True)
    rows = list(wb["IMD"].iter_rows(values_only=True))
    out = {}
    for r in rows[1:]:
        if r[0]:
            out[r[0]] = {"name": r[1], "avg_rank": r[2], "rank_avg_rank": r[3], "avg_score": r[4],
                         "rank_avg_score": r[5], "prop10": r[6], "rank_prop10": r[7]}
    return out


def load_imd_recomputed(cache):
    """IMD 2025 for the ICBs as at April 2026, from LSOA-level File 7 and the ONS LSOA to
    ICB (April 2026) lookup. Population-weighted mean IMD score (mid-2022 population) and
    the share of LSOAs in the most deprived national decile: the same two measures File 13
    publishes for the 2024 ICBs, and test_icb_pack.py reproduces File 13 on the ICBs that
    did not change. Ranks are among the ICBs returned (36)."""
    look = {}
    with open(os.path.join(cache, "lsoa21-icb26-lookup.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            look[r["LSOA21CD"]] = r["ICB26CDH"]
    agg = collections.defaultdict(lambda: [0.0, 0.0, 0, 0])
    with open(os.path.join(cache, "imd2025-file7-lsoa.csv"), encoding="utf-8-sig") as f:
        rd = csv.DictReader(f)
        cols = rd.fieldnames
        sc = next(c for c in cols if c.startswith("Index of Multiple Deprivation (IMD) Score"))
        dc = next(c for c in cols if c.startswith("Index of Multiple Deprivation (IMD) Decile"))
        pc = "Total population: mid 2022"
        for r in rd:
            icb = look.get(r["LSOA code (2021)"])
            if not icb:
                continue
            a = agg[icb]
            pop = float(r[pc])
            a[0] += pop
            a[1] += pop * float(r[sc])
            a[2] += 1
            a[3] += 1 if r[dc].strip() == "1" else 0
    out = {k: {"avg_score": v[1] / v[0], "prop10": v[3] / v[2], "lsoas": v[2], "population": v[0]}
           for k, v in agg.items()}
    for rank, k in enumerate(sorted(out, key=lambda k: -out[k]["avg_score"]), 1):
        out[k]["rank_avg_score"] = rank
    return out


def norm_icb(n):
    n = (n or "").upper().replace(" INTEGRATED CARE BOARD", "").replace(" ICB", "")
    n = re.sub(r"^NHS ", "", n)
    return re.sub(r"[^A-Z0-9]+", " ", n).strip()

def pct(a, b):
    return None if not b else 100.0 * a / b


def trend(series, floor):
    base, now = series[0], series[-1]
    if base < floor:
        return None
    return 100.0 * (now - base) / base




# --------------------------------------------------------------------------
# repo shards: ICB-level items for market and each brand word
# --------------------------------------------------------------------------
def load_shards():
    idx = json.load(open(os.path.join(GP_DIR, "index.json")))
    n = len(idx["periods"])
    items = collections.defaultdict(lambda: collections.defaultdict(lambda: [0] * n))   # icb -> brand key -> [n]
    market = collections.defaultdict(lambda: [0] * n)
    products = collections.defaultdict(lambda: collections.defaultdict(int))             # icb -> label -> 12m items
    for fn in ("s-1302.json", "s-2122.json"):
        s = json.load(open(os.path.join(GP_DIR, fn)))["s"]
        for sub, v in s.items():
            if not (sub.startswith("130201") or sub.startswith("2122")):
                continue
            for icb, prods in v["t"].items():
                for key, arr in prods.items():
                    label = v["p"].get(key, key)
                    first = label.split()[0] if label.split() else label
                    bk = brand_key(first)
                    for k in range(n):
                        market[icb][k] += arr[k] or 0
                        items[icb][bk][k] += arr[k] or 0
                    products[icb][label] += sum(arr[-12:])
    return idx, items, market, products


def cv(values):
    if len(values) < 3:
        return None
    m = statistics.mean(values)
    return None if m == 0 else statistics.pstdev(values) / m


def pretty(name):
    n = re.sub(r"^NHS ", "", name or "")
    n = re.sub(r" INTEGRATED CARE BOARD$", "", n, flags=re.I)
    n = n.title()
    return re.sub(r"\b(And|Of|The|On)\b", lambda m: m.group(1).lower(), n)


def r2(x, d=2):
    return None if x is None else round(x, d)


def assemble(cache):
    """Everything the tool needs, for every ICB and every brand. Returns (index, {icb: shard})."""
    idx, s_items, s_market, s_products = load_shards()
    periods, floor = idx["periods"], idx["minBaselineItems"]
    n = len(periods)
    last12 = range(n - 12, n)
    gmap, size = load_gpreg(cache)
    rtt, rtt_period = load_rtt(cache)
    rtt_hist = load_rtt_history(cache)
    rtt_periods = sorted(rtt_hist)
    imd = load_imd_recomputed(cache)
    prac = json.load(open(os.path.join(cache, "epd-practice-brand-months.json")))
    sources = json.load(open(os.path.join(cache, "sources.json")))

    # brand list: first words with enough national items, display name from NHSBSA
    tot, disp = collections.Counter(), {}
    for p in periods[-12:]:
        for pr, bw in prac[p].items():
            for w, (i, c) in bw.items():
                tot[brand_key(w)] += i
                disp.setdefault(brand_key(w), w)
    brands = [k for k, v in tot.most_common() if v >= BRAND_FLOOR and k not in GENERIC_WORDS]
    bset = set(brands)

    def bk_of(w):
        k = brand_key(w)
        return k if k in bset else OTHER

    # practice rows -> (icb, pcn)
    pcn = {}
    nat = {"items": collections.defaultdict(lambda: [0.0] * n), "cost": collections.defaultdict(lambda: [0.0] * n)}
    icb_cost = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0] * n))
    unmapped = 0.0
    for k, p in enumerate(periods):
        for pr, bw in prac[p].items():
            g = gmap.get(pr)
            for w, (i, c) in bw.items():
                b = bk_of(w)
                nat["items"][b][k] += i
                nat["items"]["_market"][k] += i
                nat["cost"][b][k] += c
                nat["cost"]["_market"][k] += c
                if not g:
                    unmapped += i
                    continue
                icb_cost[g["icb"]][b][k] += c
                icb_cost[g["icb"]]["_market"][k] += c
                d = pcn.setdefault((g["icb"], g["pcn"]), {"name": g["pcn_name"], "practices": set(),
                                                           "m": [0.0] * n, "b": collections.defaultdict(lambda: [0.0] * n)})
                d["practices"].add(pr)
                d["m"][k] += i
                d["b"][b][k] += i
    for d in pcn.values():
        d["list"] = sum(size.get(pr, 0) for pr in d["practices"])

    shards, rank = {}, {b: {} for b in brands}
    icb_meta = {}
    for code, name in idx["icbs"].items():
        if code == "-":
            continue
        pcns = [(c, d) for (i, c), d in pcn.items() if i == code and d["list"]]
        mkt = s_market[code]
        brand_blocks = {}
        for b in brands:
            ser = s_items[code].get(b, [0] * n)
            brand_blocks[b] = {"items": ser, "cost": [r2(x) for x in icb_cost[code][b]]}
            per = [sum(d["b"][b][k] for k in last12) / d["list"] * 1000 for _, d in pcns]
            rank[b][code] = {"items12": sum(ser[-12:]), "share12": r2(pct(sum(ser[-12:]), sum(mkt[-12:])), 3),
                             "trend": r2(trend(ser, floor), 1) if floor is not None else None,
                             "cv": r2(cv(per), 3), "nPcn": len(per)}
        shards[code] = {
            "code": code, "name": name, "periods": periods,
            "market": {"items": mkt, "cost": [r2(x) for x in icb_cost[code]["_market"]]},
            "brands": brand_blocks,
            "otherBrands": {"items": [max(0, mkt[k] - sum(s_items[code].get(b, [0] * n)[k] for b in brands)) for k in range(n)]},
            "topProducts": [{"label": l, "items12": v} for l, v in sorted(s_products[code].items(), key=lambda kv: -kv[1])[:15]],
            "pcns": [{"code": c, "name": d["name"], "list": d["list"], "practices": len(d["practices"]),
                      "market": [int(x) for x in d["m"]],
                      "brands": {b: [int(x) for x in d["b"][b]] for b in brands}}
                     for c, d in sorted(pcns, key=lambda cd: cd[1]["name"])],
        }
        series = None
        if all(code in rtt_hist[q] for q in rtt_periods):
            series = {"total": [rtt_hist[q][code]["total"] for q in rtt_periods],
                      "median": [r2(rtt_hist[q][code]["median"], 1) for q in rtt_periods],
                      "w52": [rtt_hist[q][code]["w52"] for q in rtt_periods]}
        r = rtt.get(code)
        icb_meta[code] = {
            "name": name, "pretty": pretty(name),
            "imd": ({"score": r2(imd[code]["avg_score"], 2), "rank": imd[code]["rank_avg_score"],
                     "prop10": r2(imd[code]["prop10"], 4)} if code in imd else None),
            "rtt": ({k: (r2(v, 1) if k == "median" else v) for k, v in r.items() if k != "name"} if r else None),
            "rttSeries": series,
            "list": sum(d["list"] for _, d in pcns),
        }
    index = {
        "generatedOn": datetime.date.today().isoformat(), "periods": periods, "minBaselineItems": floor,
        "rttPeriod": rtt_period, "rttPeriods": rtt_periods,
        "market": "BNF 130201 (emollient preparations) plus appliance section 2122 (emollients), primary care items",
        "brands": [{"key": b, "label": disp[b]} for b in brands],
        "icbs": icb_meta, "rank": rank,
        "national": {"items": {b: [int(x) for x in nat["items"][b]] for b in brands + ["_market"]},
                     "cost": {b: [r2(x, 0) for x in nat["cost"][b]] for b in brands + ["_market"]},
                     "rtt": rtt.get("ENGLAND"),
                     "rttSeries": {"total": [rtt_hist[q]["ENGLAND"]["total"] for q in rtt_periods],
                                   "median": [r2(rtt_hist[q]["ENGLAND"]["median"], 1) for q in rtt_periods]}},
        "unmappedItems": int(unmapped),
        "sources": {k: v for k, v in sources.items()},
        "licences": [{"source": a, "licence": b} for a, b in LICENCES],
    }
    return index, shards


NOTICE = {
    "owner": "Elevate and Thrive Ltd, company number 17154474, England and Wales",
    "copyright": "© Elevate and Thrive Ltd 2026. All rights reserved.",
    "databaseRight": "Database right asserted. Compiled from Open Government Licence sources listed in 'licences'; "
                     "the compilation, definitions and derived measures are Elevate and Thrive Ltd's.",
    "published": "Published solely to serve the Medical Sales Intelligence Hub (ICB pack tool). Public accessibility "
                 "of this file is a technical consequence of how the Hub loads it and is not a licence to use it.",
    "terms": "https://medsalesintelligencehub.co.uk/terms/",
}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--cache", default=os.path.join(ROOT, "tmp", "icb-pack-cache"))
    ap.add_argument("--out", default=os.path.join(ROOT, "data", "icb-pack"))
    ap.add_argument("--refresh", action="store_true", help="re-download every source")
    ap.add_argument("--no-fetch", action="store_true", help="use the cache as it stands")
    a = ap.parse_args(argv)
    idx = json.load(open(os.path.join(GP_DIR, "index.json")))
    if not a.no_fetch:
        fetch_all(a.cache, a.refresh)
        fetch_practice_level(a.cache, idx["periods"], a.refresh)
        fetch_rtt_history(a.cache, a.refresh)
    index, shards = assemble(a.cache)
    os.makedirs(a.out, exist_ok=True)
    for old in glob.glob(os.path.join(a.out, "icb-*.json")):
        os.remove(old)
    index["_notice"] = NOTICE
    json.dump(index, open(os.path.join(a.out, "index.json"), "w"), separators=(",", ":"))
    for code, sh in shards.items():
        sh["_notice"] = NOTICE
        json.dump(sh, open(os.path.join(a.out, "icb-%s.json" % code), "w"), separators=(",", ":"))
    size = sum(os.path.getsize(f) for f in glob.glob(os.path.join(a.out, "*.json")))
    print("wrote %d ICB shards + index to %s (%.1f MB), %d brands" % (len(shards), a.out, size / 1e6, len(index["brands"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
